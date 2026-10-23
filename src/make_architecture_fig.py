"""生成答辩用架构图（最终版）：TeaWindow × StartLux 的 1+5 可问责决策链。

输出：
  submit/30_架构图_1plus5.png   （16:9）
  submit/30_架构图_1plus5.pptx  （单页 16:9，含讲稿备注）
"""
import os

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "submit")

BG = "#0f1720"
PANEL = "#16212b"
PANEL2 = "#1d2b37"
FG = "#e8eef3"
MUTED = "#8aa0b2"
ACCENT = "#4aa8ff"
GREEN = "#3fb950"
AMBER = "#d29922"
RED = "#e5534b"
LINE = "#2b3d4d"

plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "DejaVu Sans"]
plt.rcParams["axes.unicode_minus"] = False


def box(ax, x, y, w, h, fc=PANEL, ec=LINE, r=0.08, lw=1.2):
    ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle=f"round,pad=0.02,rounding_size={r}",
                                linewidth=lw, edgecolor=ec, facecolor=fc, zorder=2))


def arrow(ax, x1, y1, x2, y2, color=MUTED, lw=1.6):
    ax.add_patch(FancyArrowPatch((x1, y1), (x2, y2), arrowstyle="-|>", mutation_scale=13,
                                 linewidth=lw, color=color, zorder=3))


def main():
    fig = plt.figure(figsize=(16, 9), dpi=190, facecolor=BG)
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, 16)
    ax.set_ylim(0, 9)
    ax.axis("off")
    ax.set_facecolor(BG)

    ax.text(0.7, 8.45, "TeaWindow × StartLux：1 + 5 可问责决策链", color=FG, fontsize=24, weight="bold")
    ax.text(0.7, 8.05, "预测数值归自研（可追溯）；StartLux 只做放行与事实核对；闭环有界、失败放行",
            color=MUTED, fontsize=12)
    ax.text(15.3, 8.45, "ONE PLACE · ONE CLIMATE PROBLEM · ONE AI SOLUTION",
            color="#3b5566", fontsize=8, ha="right")

    # 六个环节
    y0, h, w, gap = 5.35, 1.9, 2.2, 0.18
    boxes = [
        ("① 用户问题", "提问", ["今天采不采", "今晚防不防", "哪块地先浇"], MUTED),
        ("② 自研生成", "生成", ["规则 + GBDT + 模板", "数值/日期/行动", "全部来自数据"], ACCENT),
        ("③ 数据体检", "确定性", ["新鲜度 / 空洞", "上游漂移 / 刷新", "血缘 → score"], GREEN),
        ("④ StartLux", "放行 + 核对", ["supported 证据", "actionable 行动", "safe / usability"], AMBER),
        ("⑤ 有界修复", "最多 2 轮", ["仅 annotate 触发", "三条固定文案", "review 只提示"], "#b07cf0"),
        ("⑥ 输出", "交付", ["白话建议", "地块体检卡", "问答 + 审计链"], GREEN),
    ]
    xs = [0.95 + i * (w + gap) for i in range(6)]
    for i, (title, tag, lines, color) in enumerate(boxes):
        box(ax, xs[i], y0, w, h)
        ax.add_patch(FancyBboxPatch((xs[i], y0 + h - 0.09), w, 0.09,
                                    boxstyle="round,pad=0.0,rounding_size=0.04",
                                    linewidth=0, facecolor=color, alpha=0.85, zorder=3))
        ax.text(xs[i] + 0.14, y0 + h - 0.42, title, color=FG, fontsize=12.5, weight="bold", zorder=4)
        ax.text(xs[i] + 0.14, y0 + h - 0.72, tag, color=color, fontsize=8.8, zorder=4)
        for j, ln in enumerate(lines):
            ax.text(xs[i] + 0.14, y0 + 0.88 - j * 0.27, ln, color="#c6d5e2", fontsize=9.2, zorder=4)
    for i in range(5):
        arrow(ax, xs[i] + w + 0.01, y0 + h / 2, xs[i + 1] - 0.01, y0 + h / 2, color="#5a7184")

    # 判定三态
    ax.text(0.95, 4.82, "判定三态（都不阻断主流程）：", color=MUTED, fontsize=9.8)
    chips = [("ok", GREEN, "证据一致 · 直接给结论"),
             ("review", AMBER, "置信一般 · 提示实勘"),
             ("annotate", RED, "证据不足 · 先核验")]
    cx = 0.95
    for name, col, desc in chips:
        box(ax, cx, 4.25, 0.85, 0.38, fc=PANEL2, ec=col, r=0.06)
        ax.text(cx + 0.425, 4.44, name, color=col, fontsize=9.5, ha="center", weight="bold")
        ax.text(cx + 0.95, 4.44, desc, color="#c6d5e2", fontsize=8.8, va="center")
        cx += 0.85 + 1.45

    # 左：边界
    box(ax, 0.7, 1.45, 7.55, 2.55)
    ax.text(0.95, 3.72, "边界、失败策略与实验结论", color=FG, fontsize=11.5, weight="bold")
    left = [
        ("fail-open", "未配置 / 超时 10s / 异常 → 原样返回自研答案"),
        ("不改数", "模型只能加提示，不能改数值、日期与行动"),
        ("不越权", "人工实勘与官方预警拥有最终决定权"),
        ("可追溯", "来源链 + 证据快照 + 版本 SHA-256，可在 /qa_review 复核"),
        ("实验结论", "霜型分组订正 1033 天验证：MAE 1.00→1.03，不上线，只做元信息"),
    ]
    for i, (k, v) in enumerate(left):
        yy = 3.33 - i * 0.38
        ax.text(0.98, yy, k, color=ACCENT, fontsize=9.5, weight="bold")
        ax.text(1.98, yy, v, color="#c6d5e2", fontsize=8.9)

    # 右：预测辅助 + OpenCode
    box(ax, 8.45, 1.45, 6.85, 2.55)
    ax.text(8.7, 3.72, "预测辅助 · Agent 复用", color=FG, fontsize=11.5, weight="bold")
    ax.text(8.7, 3.33, "GET /api/forecast/assist", color=AMBER, fontsize=9.6, weight="bold")
    ax.text(11.35, 3.33, "霜型 / 置信 / 一致性（缓存 0.4s，不改数值）", color="#c6d5e2", fontsize=8.9)
    ax.text(8.7, 2.94, "OpenCode MCP（stdio）", color=AMBER, fontsize=9.6, weight="bold")
    ax.text(11.35, 2.94, "startlux_decide · startlux_choose", color="#c6d5e2", fontsize=8.9)
    ax.text(8.7, 2.55, "key 从 .env 读取，不写入配置；本地起服务即用", color=MUTED, fontsize=8.6)
    ax.text(8.7, 2.14, "问答链路：生成不交给大模型，判断才交给它", color="#c6d5e2", fontsize=8.9)
    ax.text(8.7, 1.78, "值班口径：规则为主，模型举手，人拍板", color=MUTED, fontsize=8.6)

    # 指标条
    box(ax, 0.7, 0.42, 14.6, 0.85, fc=PANEL2)
    metrics = [
        ("AUC 0.998 / F1 0.85", "冻结评估集 n=1014"),
        ("官方开采日 MAE 1.8 天", "前向验证 2.67 天"),
        ("155 tests", "fail-open 覆盖"),
        ("护栏 1.5–2.8s", "超时即放行"),
    ]
    mx = 1.0
    for big, small in metrics:
        ax.text(mx, 0.94, big, color=FG, fontsize=10.8, weight="bold")
        ax.text(mx, 0.62, small, color=MUTED, fontsize=8.6)
        mx += 3.62

    ax.text(0.7, 0.15, "口径：弱标签 + ERA5 对照，非田间产量真值；建议仅作决策辅助，不替代实地核查与官方预警。",
            color="#5c7183", fontsize=8)

    png = os.path.join(OUT, "30_架构图_1plus5.png")
    fig.savefig(png, facecolor=BG, bbox_inches="tight", pad_inches=0.16)
    plt.close(fig)
    print("->", png, os.path.getsize(png) // 1024, "KB")

    try:
        from pptx import Presentation
        from pptx.util import Inches
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(png, 0, 0, width=prs.slide_width)
        slide.notes_slide.notes_text_frame.text = (
            "60 秒讲解：\n"
            "1) 六个环节：问题 → 自研生成 → 数据体检 → StartLux 放行/核对 → 有界修复 → 输出。生成不交给大模型。\n"
            "2) 三种判定：ok / review / annotate，都不阻断主流程；只有 annotate 才进入修复，最多 2 轮。\n"
            "3) 左下边界：fail-open、不改数、不越权、可追溯；霜型分组实验 1033 天验证无收益，主动不上线。\n"
            "4) 右下：预测辅助只给元信息；同一个模型在 OpenCode 里作为 MCP 工具复用。\n"
            "5) 底部数字：AUC 0.998 / F1 0.85（冻结评估集）；官方开采日 MAE 1.8 天、前向 2.67 天。"
        )
        pptx = os.path.join(OUT, "30_架构图_1plus5.pptx")
        prs.save(pptx)
        print("->", pptx, os.path.getsize(pptx) // 1024, "KB")
    except Exception as e:  # noqa: BLE001
        print("pptx skip:", e)


if __name__ == "__main__":
    main()
