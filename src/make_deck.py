"""
生成路演幻灯片（16:9）：submit/TeaWindow_路演.pptx（中文）与 submit/TeaWindow_Pitch_EN.pptx（英文）。
依赖：python-pptx。
"""
import os

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIG = os.path.join(ROOT, "report", "figures")
ASSETS = os.path.join(ROOT, "web", "demo_assets")
OUT_DIR = os.path.join(ROOT, "submit")

DARK = RGBColor(0x0B, 0x12, 0x19)
FG = RGBColor(0xE8, 0xEE, 0xF3)
MUTED = RGBColor(0x9F, 0xB4, 0xC6)
ACCENT = RGBColor(0x4A, 0xA8, 0xFF)

CN_FONT = "Microsoft YaHei"
EN_FONT = "Segoe UI"


def _img(name):
    for base in (ASSETS, FIG):
        p = os.path.join(base, name)
        if os.path.exists(p):
            return p
    return None


CN = [
    {"t": "梅坞采候 TeaWindow", "sub": "AI4Climate 赛道 01 · 气候智能茶园",
     "b": ["面向杭州梅家坞西湖龙井核心产区（30.21°N, 120.10°E）",
           "把「倒春寒 × 高温 × 干旱」变成可执行的地块级判断",
           "One Place · One Climate Problem · One AI Solution"], "cover": True},
    {"t": "问题与证据：真实且可核验", "img": "fig_yearly.png",
     "b": ["约 70% 茶园为早芽、抗旱弱的龙井43号",
           "2010 浙江冻害：约 150 万亩、损失约 17 亿元（西湖区近 3 成受冻、推迟约 10 天）",
           "2020/2022 西湖区约 5000 亩高温灼伤；2024 杭州高温 60 天（破纪录）"]},
    {"t": "用户与痛点", "b": ["用户：茶农 / 采茶队 / 合作社 / 茶企",
                              "每天三问：今天有风险吗？是什么风险？该采还是该防？",
                              "痛点：气象·遥感·土壤·地形信号分散，难以落到地块级行动",
                              "边界：决策辅助，需人工复核、以官方预警为准"]},
    {"t": "方案架构：把信号变成行动", "img": "dashboard.png",
     "b": ["输入：Open-Meteo ERA5 历史 + MET Norway 预报（双源互备）· NDVI · 地形",
           "指标引擎：积温/霜冻/倒春寒/高温/干旱/连阴雨",
           "模型：numpy GBDT 复合风险 → 透明规则采摘适宜度与最佳窗口",
           "输出：茶农白话行动建议（离线/在线双形态）"]},
    {"t": "AI 为什么必要", "img": "fig_importance.png",
     "b": ["单阈值（如仅 Tmin）无法刻画「霜冻×高温×干旱×物候」交互",
           "留出年 2024–2025：AUC 0.998 / F1 0.85 vs 单阈值 0.78 / 0.50（弱标签复现）",
           "集合预报给不确定性区间；地形 + NDVI 做地块级解释",
           "明确：弱标签 ≠ 田间精度"]},
    {"t": "从预测到决策（新颖点）", "img": "fig_insurance.png",
     "b": ["采摘调度排班 + 决策论期望效用（抢采 / 等待 / 防护）",
           "价值闭环：预期损失 / 保费；天气指数保险（地块级定价）",
           "地块级霜冻易感（地形/坡向）"]},
    {"t": "在线演示 + 自研问答", "img": "chat.png",
     "b": ["实时 https://teawindow.onrender.com/ ；静态 https://eric2007-sudo.github.io/TeaWindow/",
           "自研问答：ONNX 意图 + 事实接地 + 来源链",
           "多维复合提问分项作答 · JSON/表格跨轮一致 · 任务引用与记忆",
           "运行闭环：3h 刷新 + 漂移检测 + 过期降级 + 告警（钉钉加签真实送达）"]},
    {"t": "验证与路线图", "img": "fig_backtest.png",
     "b": ["春霜 357 天：Tmin MAE 1.32℃(CI 1.21–1.43) vs 持续法 2.08℃",
           "1/3/7 天 MAE 0.49/0.67/0.91℃；POD 0.709 / FAR 0 / CSI 0.709",
           "事件回测 7/7；下一步：冻结版本 + 地块级传感器 + 正/负田间记录",
           "90 天真实场景验证计划；决策辅助，不替代官方预警"]},
    {"t": "如何运行与合规", "b": ["在线：https://teawindow.onrender.com/ ；静态：https://eric2007-sudo.github.io/TeaWindow/",
                     "本地：uvicorn src.serve:app ；离线单页：web/index.html 双击即开",
                     "数据：Open-Meteo ERA5 历史 + MET Norway 预报（双源）· NASA GIBS NDVI · Sentinel-2 · DEM",
                     "许可：代码 MIT · 数据/文档 CC BY 4.0；决策辅助，不替代官方预警"]},
    {"t": "ONE PLACE · ONE CLIMATE PROBLEM · ONE AI SOLUTION",
     "sub": "梅坞采候 TeaWindow ｜ 让 AI 走进真实茶园",
     "b": ["在线 https://teawindow.onrender.com/",
           "静态 https://eric2007-sudo.github.io/TeaWindow/",
           "仓库 https://github.com/Eric2007-sudo/TeaWindow"], "cover": True},
]

EN = [
    {"t": "TeaWindow", "sub": "AI4Climate · Climate-smart tea farming",
     "b": ["Meijiawu, West Lake Longjing core area (30.21N, 120.10E)",
           "Turning late frost x heat x drought into plot-level action",
           "One Place - One Climate Problem - One AI Solution"], "cover": True},
    {"t": "Problem & evidence", "img": "fig_yearly.png",
     "b": ["~70% of gardens grow early-budding, drought-sensitive Longjing 43",
           "2010 Zhejiang cold event: ~1.5M mu, ~RMB 1.7bn losses (~30% of West Lake gardens hit, ~10-day delay)",
           "2020/2022: ~5,000 mu heat damage; 2024 Hangzhou 60 high-temp days (record)"]},
    {"t": "Users & pain point", "b": ["Farmers, harvesting teams, cooperatives, tea enterprises",
                                      "Three daily questions: risk today? which risk? harvest or protect?",
                                      "Signals (weather/vegetation/terrain/soil) are scattered",
                                      "Decision support only; respects official warnings"]},
    {"t": "Solution architecture", "img": "dashboard.png",
     "b": ["Input: Open-Meteo ERA5 history + MET Norway forecast (dual-source), NDVI, terrain",
           "Indicator engine: GDD, frost, cold spells, heat, drought, prolonged rain",
           "Model: numpy GBDT compound risk -> transparent rules for suitability & windows",
           "Output: plain-language actions (offline/online)"]},
    {"t": "Why AI is necessary", "img": "fig_importance.png",
     "b": ["A single threshold (e.g. Tmin) misses frost x heat x drought x stage interactions",
           "Held-out 2024-2025: AUC 0.998 / F1 0.85 vs single-threshold 0.78 / 0.50 (weak-label)",
           "Ensemble gives uncertainty; terrain + NDVI give plot-level context",
           "Weak-label evaluation is not field accuracy"]},
    {"t": "From prediction to decisions", "img": "fig_insurance.png",
     "b": ["Harvest scheduling + decision-theoretic expected utility (pick/wait/protect)",
           "Value loop: expected loss / premium; weather-index insurance (plot-level)",
           "Plot-level frost susceptibility from terrain"]},
    {"t": "Live demo + in-house QA", "img": "chat.png",
     "b": ["Live https://teawindow.onrender.com/ ; Static https://eric2007-sudo.github.io/TeaWindow/",
           "In-house QA: ONNX intent + grounded generation + provenance chain",
           "Multi-part questions, JSON/table cross-turn consistency, task reference & memory",
           "Ops loop: 3h refresh + drift check + stale degradation + real DingTalk alerts"]},
    {"t": "Validation & roadmap", "img": "fig_backtest.png",
     "b": ["Spring frost 357 days: Tmin MAE 1.32C (CI 1.21-1.43) vs 2.08C persistence",
           "1/3/7-day MAE 0.49/0.67/0.91C; POD 0.709 / FAR 0 / CSI 0.709",
           "Event backtest 7/7; next: frozen versions + plot sensors + positive/negative field reports",
           "90-day field pilot; decision support only"]},
    {"t": "Run & compliance", "b": ["Live: https://teawindow.onrender.com/ ; Static: https://eric2007-sudo.github.io/TeaWindow/",
                                "Local: uvicorn src.serve:app ; Offline single-page: web/index.html",
                                "Data: Open-Meteo ERA5 history + MET Norway forecast (dual-source), NASA GIBS NDVI, Sentinel-2, DEM",
                                "License: code MIT, data/docs CC BY 4.0; decision support only"]},
    {"t": "ONE PLACE - ONE CLIMATE PROBLEM - ONE AI SOLUTION",
     "sub": "TeaWindow - bringing AI to real tea gardens",
     "b": ["Live https://teawindow.onrender.com/",
           "Static https://eric2007-sudo.github.io/TeaWindow/",
           "Repo https://github.com/Eric2007-sudo/TeaWindow"], "cover": True},
]


NOTES_CN = [
    "开场：一杯西湖龙井取决于短短几天的采摘窗口，而气候正在打乱它。我们做的是让信号变成地块级行动。",
    "证据来自公开报道与官方统计：2010 冻害 17 亿、西湖区近 3 成受冻；2020/2022 约 5000 亩灼伤；2024 杭州高温 60 天。强调“问题真实、可核验”。",
    "用户每天只问三件事：今天有风险吗、是什么风险、该采还是该防。痛点是信号分散，难落到地块。边界：决策辅助、以官方预警为准。",
    "讲清数据双源（Open-Meteo 历史 + MET Norway 预报，限流自动回退）、指标引擎、GBDT 复合风险、透明规则与白话建议；强调“LLM 只润色、规则兜底”。",
    "讲清为什么单阈值不够：霜冻×高温×干旱×物候交互。给出留出年 AUC 0.998/F1 0.85 vs 单阈值 0.78/0.50，并明确这是弱标签复现、不等于田间精度。",
    "从预测走到决策：采摘排班 + 期望效用（抢采/等待/防护）+ 价值闭环与天气指数保险 + 地块级霜冻易感。",
    "现场演示：打开在线 Demo，点顶部“智能问答”提问，展示多维复合提问分项作答、JSON/表格跨轮一致；再点“扫码分享”看二维码。",
    "验证：春霜 357 天 MAE 1.32℃(CI 1.21–1.43) vs 持续法 2.08℃；POD 0.709/FAR 0/CSI 0.709；事件回测 7/7。下一步冻结版本+地块级传感器+正负田间记录。",
    "运行与合规：在线/静态双形态；本地 uvicorn 或离线单页双击即开；数据全部公开免密钥；代码 MIT、数据文档 CC BY 4.0。",
    "收尾：One Place · One Climate Problem · One AI Solution。给出在线/静态/仓库三个地址。",
]
NOTES_EN = [
    "Hook: a good cup of West Lake Longjing depends on a few days of harvest window, and climate is disrupting it.",
    "Evidence is public and citable: 2010 losses ~RMB 1.7bn; 2020/2022 ~5,000 mu heat damage; 2024 Hangzhou 60 high-temp days.",
    "Three daily questions: risk today? which risk? harvest or protect? Signals are scattered; decision support only.",
    "Dual-source weather (Open-Meteo history + MET Norway forecast with fallback), indicator engine, GBDT compound risk, transparent rules; LLM is optional rephrasing with rule fallback.",
    "One threshold cannot capture frost x heat x drought x stage interactions. Held-out AUC 0.998/F1 0.85 vs 0.78/0.50; explicitly weak-label, not field accuracy.",
    "From prediction to decisions: scheduling + expected utility (pick/wait/protect) + value/insurance + plot-level frost susceptibility.",
    "Live demo: open the site, click In-house QA, ask a multi-part question, then show the QR share page.",
    "Validation: 357-day spring frost MAE 1.32C vs 2.08C persistence; POD 0.709/FAR 0/CSI 0.709; 7/7 backtest. Next: frozen versions + plot sensors + field reports.",
    "Run & compliance: live/static forms; local uvicorn or offline single page; all data public and key-free; code MIT, data/docs CC BY 4.0.",
    "Close with the tagline and the three links (live/static/repo).",
]


def build(slides, font, out_name, notes=None):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for si, s in enumerate(slides):
        sl = prs.slides.add_slide(blank)
        if notes and si < len(notes):
            try:
                sl.notes_slide.notes_text_frame.text = notes[si]
            except Exception:  # noqa: BLE001
                pass
        bg = sl.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid(); bg.fill.fore_color.rgb = DARK; bg.line.fill.background()
        # 顶栏
        bar = sl.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.12))
        bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background()

        if s.get("cover"):
            tb = sl.shapes.add_textbox(Inches(0.9), Inches(2.2), Inches(11.5), Inches(1.6))
            p = tb.text_frame.paragraphs[0]; r = p.add_run(); r.text = s["t"]
            r.font.size = Pt(40); r.font.bold = True; r.font.color.rgb = FG; r.font.name = font
            if s.get("sub"):
                sb = sl.shapes.add_textbox(Inches(0.9), Inches(3.8), Inches(11.5), Inches(0.7))
                rr = sb.text_frame.paragraphs[0].add_run(); rr.text = s["sub"]
                rr.font.size = Pt(20); rr.font.color.rgb = MUTED; rr.font.name = font
            bb = sl.shapes.add_textbox(Inches(0.9), Inches(4.7), Inches(11.5), Inches(1.6))
            tf = bb.text_frame; tf.word_wrap = True
            for i, line in enumerate(s["b"]):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                rr = p.add_run(); rr.text = "• " + line
                rr.font.size = Pt(15); rr.font.color.rgb = MUTED; rr.font.name = font
            continue

        tb = sl.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(12), Inches(1.0))
        r = tb.text_frame.paragraphs[0].add_run(); r.text = s["t"]
        r.font.size = Pt(28); r.font.bold = True; r.font.color.rgb = FG; r.font.name = font

        img = _img(s["img"]) if s.get("img") else None
        bw = Inches(6.6) if img else Inches(11.8)
        bb = sl.shapes.add_textbox(Inches(0.7), Inches(1.7), bw, Inches(5.0))
        tf = bb.text_frame; tf.word_wrap = True
        for i, line in enumerate(s["b"]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(8)
            rr = p.add_run(); rr.text = "• " + line
            rr.font.size = Pt(16); rr.font.color.rgb = FG; rr.font.name = font
        if img:
            try:
                sl.shapes.add_picture(img, Inches(7.5), Inches(1.8), height=Inches(4.4))
            except Exception:  # noqa: BLE001
                pass
    os.makedirs(OUT_DIR, exist_ok=True)
    out = os.path.join(OUT_DIR, out_name)
    prs.save(out)
    print("saved", out, round(os.path.getsize(out) / 1e6, 2), "MB")


if __name__ == "__main__":
    build(CN, CN_FONT, "TeaWindow_路演.pptx", NOTES_CN)
    build(EN, EN_FONT, "TeaWindow_Pitch_EN.pptx", NOTES_EN)