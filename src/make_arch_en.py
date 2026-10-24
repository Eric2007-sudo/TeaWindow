"""English version of the 1+5 architecture figure (for foreign judges)."""
import os

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "submit")

BG, PANEL, PANEL2 = "#0f1720", "#16212b", "#1d2b37"
FG, MUTED, LINE = "#e8eef3", "#8aa0b2", "#2b3d4d"
ACCENT, GREEN, AMBER, RED = "#4aa8ff", "#3fb950", "#d29922", "#e5534b"

plt.rcParams["font.sans-serif"] = ["Segoe UI", "Microsoft YaHei", "DejaVu Sans"]
plt.rcParams["axes.unicode_minus"] = False


def box(ax, x, y, w, h, fc=PANEL, ec=LINE, r=0.08):
    ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle=f"round,pad=0.02,rounding_size={r}",
                                linewidth=1.2, edgecolor=ec, facecolor=fc, zorder=2))


def arrow(ax, x1, y1, x2, y2):
    ax.add_patch(FancyArrowPatch((x1, y1), (x2, y2), arrowstyle="-|>", mutation_scale=13,
                                 linewidth=1.6, color=MUTED, zorder=3))


def main():
    fig = plt.figure(figsize=(16, 9), dpi=190, facecolor=BG)
    ax = fig.add_axes([0, 0, 1, 1]); ax.set_xlim(0, 16); ax.set_ylim(0, 9)
    ax.axis("off"); ax.set_facecolor(BG)

    ax.text(0.7, 8.45, "TeaWindow x StartLux: the 1 + 5 accountable decision chain", color=FG,
            fontsize=22, weight="bold")
    ax.text(0.7, 8.05, "Numbers come from our own pipeline; StartLux only gates and fact-checks; the loop is bounded and fails open",
            color=MUTED, fontsize=11)
    ax.text(15.3, 8.45, "ONE PLACE · ONE CLIMATE PROBLEM · ONE AI SOLUTION", color="#3b5566",
            fontsize=8, ha="right")

    y0, h, w, gap = 5.35, 1.9, 2.2, 0.18
    boxes = [
        ("1. User question", "ask", ["Harvest today?", "Protect tonight?", "Which plot first?"], MUTED),
        ("2. Our pipeline", "generate", ["Rules + GBDT + templates", "Numbers and dates", "all from data"], ACCENT),
        ("3. Data health", "deterministic", ["Freshness / gaps", "Upstream drift", "Lineage -> score"], GREEN),
        ("4. StartLux", "gate + check", ["supported: evidence", "actionable: action", "safe / usability"], AMBER),
        ("5. Bounded repair", "max 2 rounds", ["only on 'annotate'", "three fixed notes", "'review' stops"], "#b07cf0"),
        ("6. Output", "deliver", ["Plain advice", "Plot report card", "Q&A + audit chain"], GREEN),
    ]
    xs = [0.95 + i * (w + gap) for i in range(6)]
    for i, (t, tag, lines, col) in enumerate(boxes):
        box(ax, xs[i], y0, w, h)
        ax.add_patch(FancyBboxPatch((xs[i], y0 + h - 0.09), w, 0.09,
                                    boxstyle="round,pad=0.0,rounding_size=0.04",
                                    linewidth=0, facecolor=col, alpha=0.85, zorder=3))
        ax.text(xs[i] + 0.14, y0 + h - 0.42, t, color=FG, fontsize=12.2, weight="bold", zorder=4)
        ax.text(xs[i] + 0.14, y0 + h - 0.72, tag, color=col, fontsize=8.8, zorder=4)
        for j, ln in enumerate(lines):
            ax.text(xs[i] + 0.14, y0 + 0.88 - j * 0.27, ln, color="#c6d5e2", fontsize=9.2, zorder=4)
    for i in range(5):
        arrow(ax, xs[i] + w + 0.01, y0 + h / 2, xs[i + 1] - 0.01, y0 + h / 2)

    ax.text(0.95, 4.82, "Three verdicts, none blocks the answer:", color=MUTED, fontsize=9.8)
    chips = [("ok", GREEN, "evidence agrees - answer directly"),
             ("review", AMBER, "medium confidence - advise a field check"),
             ("annotate", RED, "weak evidence - verify first")]
    cx = 0.95
    for name, col, desc in chips:
        box(ax, cx, 4.25, 4.35, 0.38, fc=PANEL2, ec=col, r=0.06)
        ax.text(cx + 0.14, 4.44, name, color=col, fontsize=9.3, weight="bold", va="center")
        ax.text(cx + 0.95, 4.44, desc, color="#c6d5e2", fontsize=8.4, va="center")
        cx += 4.35 + 0.35

    # Left panel
    box(ax, 0.7, 1.45, 7.55, 2.55)
    ax.text(0.95, 3.72, "Boundaries, failure policy, and what we tested", color=FG, fontsize=11.5, weight="bold")
    left = [
        ("fail-open", "no key / 10 s timeout / error -> return our own answer"),
        ("no rewriting", "the model may add a note, never change numbers or actions"),
        ("no override", "field inspection and official warnings keep the final say"),
        ("traceable", "source chain + evidence snapshot + version SHA-256 in /qa_review"),
        ("negative result", "frost-type correction tested on 1,033 days: MAE 1.00 -> 1.03"),
    ]
    for i, (k, v) in enumerate(left):
        yy = 3.33 - i * 0.38
        ax.text(0.98, yy, k, color=ACCENT, fontsize=9.5, weight="bold")
        ax.text(2.05, yy, v, color="#c6d5e2", fontsize=8.9)

    # Right panel
    box(ax, 8.45, 1.45, 6.85, 2.55)
    ax.text(8.7, 3.72, "Prediction assist and agent reuse", color=FG, fontsize=11.5, weight="bold")
    ax.text(8.7, 3.33, "GET /api/forecast/assist", color=AMBER, fontsize=9.6, weight="bold")
    ax.text(11.6, 3.33, "frost type / confidence / consistency (cached, no numeric change)", color="#c6d5e2", fontsize=8.9)
    ax.text(8.7, 2.94, "OpenCode MCP (stdio)", color=AMBER, fontsize=9.6, weight="bold")
    ax.text(11.6, 2.94, "startlux_decide · startlux_choose", color="#c6d5e2", fontsize=8.9)
    ax.text(8.7, 2.55, "Key read from .env, never stored in config", color=MUTED, fontsize=8.6)
    ax.text(8.7, 2.14, "Generation stays in-house; only judgement is delegated", color="#c6d5e2", fontsize=8.9)
    ax.text(8.7, 1.78, "Rules lead, the model raises a hand, humans decide", color=MUTED, fontsize=8.6)

    box(ax, 0.7, 0.42, 14.6, 0.85, fc=PANEL2)
    metrics = [
        ("AUC 0.998 / F1 0.85", "frozen eval set n=1,014"),
        ("Picking-date MAE 1.8 d", "forward test 2.67 d"),
        ("155 tests", "fail-open covered"),
        ("Gate 1.5-2.8 s", "timeout returns the answer"),
    ]
    mx = 1.0
    for big, small in metrics:
        ax.text(mx, 0.94, big, color=FG, fontsize=10.8, weight="bold")
        ax.text(mx, 0.62, small, color=MUTED, fontsize=8.6)
        mx += 3.62

    ax.text(0.7, 0.15, "Weak labels with ERA5 comparison, not field yield truth; advice is decision support, not a substitute for inspection or official warnings.",
            color="#5c7183", fontsize=8)

    png = os.path.join(OUT, "30_architecture_1plus5_EN.png")
    fig.savefig(png, facecolor=BG, bbox_inches="tight", pad_inches=0.16)
    plt.close(fig)
    print("->", png, os.path.getsize(png) // 1024, "KB")

    try:
        from pptx import Presentation
        from pptx.util import Inches
        prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.shapes.add_picture(png, 0, 0, width=prs.slide_width)
        s.notes_slide.notes_text_frame.text = (
            "60-second script: six stages; generation stays in-house, StartLux only gates and fact-checks. "
            "Three verdicts never block the answer; only 'annotate' triggers a repair, at most twice. "
            "Boundaries: fail-open, no rewriting, no override, traceable; the frost-type correction was tested and rejected (1.00 to 1.03)."
        )
        p = os.path.join(OUT, "30_architecture_1plus5_EN.pptx")
        prs.save(p)
        print("->", p, os.path.getsize(p) // 1024, "KB")
    except Exception as e:  # noqa: BLE001
        print("pptx skip:", e)


if __name__ == "__main__":
    main()
